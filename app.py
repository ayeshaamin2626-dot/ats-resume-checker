"""ATS Resume Checker - Streamlit + Google Gemini Flash."""

import json
import os
import re
from io import BytesIO

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-2.5-flash"
MAX_RESUME_CHARS = 30000
MIN_RESUME_CHARS = 100

st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="centered")


# ---------------------------------------------------------------- config
def get_api_key() -> str | None:
    """Read the Gemini key from Streamlit secrets, then env var."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
        if key:
            return key
    except Exception:
        pass
    return os.getenv("GEMINI_API_KEY")


def get_model_name() -> str:
    try:
        name = st.secrets.get("GEMINI_MODEL")
        if name:
            return name
    except Exception:
        pass
    return os.getenv("GEMINI_MODEL", DEFAULT_MODEL)


# ---------------------------------------------------------------- parsing
def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")


def parse_json_response(text: str) -> dict:
    """Parse JSON from the model, tolerating code fences or extra text."""
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI returned an unreadable response. Please try again.")


def clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


# ---------------------------------------------------------------- AI
PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.

Analyze the RESUME below{jd_clause}. Be honest, specific and critical. Do not inflate scores.

Return ONLY a JSON object with exactly this structure:
{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {{
    "formatting_and_structure": <0-100>,
    "keywords_and_skills": <0-100>,
    "work_experience_impact": <0-100>,
    "education_and_certifications": <0-100>,
    "clarity_and_grammar": <0-100>
  }},
  "strengths": ["..."],
  "weaknesses": ["..."],
  "missing_keywords": ["..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "<resume section>", "issue": "<what is wrong>", "suggestion": "<exactly how to fix it, with a short example where useful>"}}
  ]
}}

Scoring guide: formatting/parsability, relevant keywords, quantified achievements, strong action verbs,
standard section headings (Summary, Experience, Education, Skills), contact info, length and consistency.
{jd_rule}
Give 3-6 strengths, 3-6 weaknesses and 5-10 improvements ordered by priority.

{jd_block}RESUME:
\"\"\"
{resume}
\"\"\"
"""


def analyze_resume(resume_text: str, job_description: str, api_key: str, model: str) -> dict:
    jd = job_description.strip()
    prompt = PROMPT.format(
        jd_clause=" against the JOB DESCRIPTION provided" if jd else "",
        jd_rule=(
            "Judge keyword match and relevance against the job description, and list important "
            "keywords from it that are missing from the resume."
            if jd
            else "No job description was given, so judge against general best practices and list "
            "commonly expected keywords for the candidate's apparent field."
        ),
        jd_block=f'JOB DESCRIPTION:\n"""\n{jd[:10000]}\n"""\n\n' if jd else "",
        resume=resume_text[:MAX_RESUME_CHARS],
    )
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    data = parse_json_response(response.text)

    # normalise so the UI never crashes on a missing field
    data["ats_score"] = clamp_score(data.get("ats_score"))
    sections = data.get("section_scores") or {}
    data["section_scores"] = {k: clamp_score(v) for k, v in sections.items()}
    for key in ("strengths", "weaknesses", "missing_keywords", "improvements"):
        if not isinstance(data.get(key), list):
            data[key] = []
    data["summary"] = str(data.get("summary", ""))
    return data


# ---------------------------------------------------------------- UI helpers
def score_label(score: int) -> str:
    if score >= 80:
        return "🟢 Excellent"
    if score >= 60:
        return "🟡 Good, needs polish"
    if score >= 40:
        return "🟠 Needs work"
    return "🔴 Poor"


def render_results(data: dict) -> None:
    score = data["ats_score"]
    st.divider()
    st.subheader("Your ATS Score")
    col1, col2 = st.columns([1, 2])
    col1.metric("Score", f"{score}/100")
    col2.write(score_label(score))
    col2.progress(score / 100)
    if data["summary"]:
        st.info(data["summary"])

    if data["section_scores"]:
        st.subheader("Score breakdown")
        for name, val in data["section_scores"].items():
            st.write(name.replace("_", " ").title())
            st.progress(val / 100, text=f"{val}/100")

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for item in data["strengths"]:
            st.markdown(f"- {item}")
    with right:
        st.subheader("⚠️ Weaknesses")
        for item in data["weaknesses"]:
            st.markdown(f"- {item}")

    if data["missing_keywords"]:
        st.subheader("🔑 Missing keywords")
        st.write(", ".join(f"`{k}`" for k in data["missing_keywords"]))

    st.subheader("🛠️ Suggested improvements")
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for imp in data["improvements"]:
        if not isinstance(imp, dict):
            st.markdown(f"- {imp}")
            continue
        pr = str(imp.get("priority", "Medium"))
        title = f"{icons.get(pr.lower(), '🟠')} {pr} - {imp.get('section', 'General')}"
        with st.expander(title):
            st.markdown(f"**Issue:** {imp.get('issue', '')}")
            st.markdown(f"**Fix:** {imp.get('suggestion', '')}")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(data, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# ---------------------------------------------------------------- main
def main() -> None:
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and concrete ways to improve it.")

    api_key = get_api_key()
    if not api_key:
        st.error(
            "Gemini API key not found. Add `GEMINI_API_KEY` to Streamlit secrets "
            "(or set it as an environment variable). See the README."
        )
        st.stop()

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, for a targeted score)",
        height=150,
        placeholder="Paste the job description here to check keyword match...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded.name, uploaded.getvalue())
            if len(text) < MIN_RESUME_CHARS:
                st.error(
                    "Couldn't read enough text. If your PDF is a scanned image, "
                    "ATS systems can't read it either - export a text-based PDF."
                )
                st.stop()
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(text, job_description, api_key, get_model_name())
            st.session_state["result"] = result
        except ValueError as e:
            st.error(str(e))
        except Exception as e:
            st.error(f"Something went wrong: {e}")

    if "result" in st.session_state:
        render_results(st.session_state["result"])

    st.caption("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")


if __name__ == "__main__":
    main()
